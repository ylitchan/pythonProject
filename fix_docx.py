from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

# 创建新文档
doc = Document()

# 添加原文
p1 = doc.add_paragraph()
run1 = p1.add_run("原文（有语病）：")
run1.bold = True
run1.font.size = Pt(14)
run1.font.color.rgb = RGBColor(255, 0, 0)

doc.add_paragraph("大概答复的水电费都是")

# 添加分隔线
doc.add_paragraph("=" * 50)

# 添加修改建议标题
p2 = doc.add_paragraph()
run2 = p2.add_run("修改建议：")
run2.bold = True
run2.font.size = Pt(14)
run2.font.color.rgb = RGBColor(0, 128, 0)

doc.add_paragraph()

# 方案1
p3 = doc.add_paragraph()
run3 = p3.add_run('方案1（如果想表达"答复内容"）：')
run3.bold = True
run3.font.size = Pt(12)
doc.add_paragraph("答复的内容大概都是关于水电费的。")
doc.add_paragraph()

# 方案2
p4 = doc.add_paragraph()
run4 = p4.add_run('方案2（如果想表达"水电费金额"）：')
run4.bold = True
run4.font.size = Pt(12)
doc.add_paragraph("水电费大概都是这些金额。")
doc.add_paragraph()

# 方案3
p5 = doc.add_paragraph()
run5 = p5.add_run('方案3（如果想表达"答复说明"）：')
run5.bold = True
run5.font.size = Pt(12)
doc.add_paragraph("关于水电费的答复大概都是这样的。")
doc.add_paragraph()

# 方案4
p6 = doc.add_paragraph()
run6 = p6.add_run('方案4（如果想表达"费用构成"）：')
run6.bold = True
run6.font.size = Pt(12)
doc.add_paragraph("大概的水电费都是由这些部分组成的。")

# 保存修改后的文档
output_file = "新建 DOCX 文档_已修改.docx"
doc.save(output_file)
print(f"✓ 已创建修改后的文档：{output_file}")
print("\n文档包含4个修改方案，请查看并选择最合适的表达方式。")
